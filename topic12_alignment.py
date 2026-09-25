from openpyxl import load_workbook
from openpyxl.styles import Alignment

wb = load_workbook("company_data.xlsx")
sheet = wb["Employees"]

# Center align all data
for row in sheet.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

# Wrap the header text
for cell in sheet[1]:
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

wb.save("company_data.xlsx")

print("Data aligned properly")