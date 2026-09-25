from openpyxl import load_workbook
from openpyxl.styles import PatternFill

wb = load_workbook("company_data.xlsx")
sheet = wb["Employees"]

# Create yellow background
yellow = PatternFill(
    fill_type="solid",
    fgColor="FFFF00"
)

# Apply yellow color to header row
for cell in sheet[1]:
    cell.fill = yellow

wb.save("company_data.xlsx")

print("Background color applied successfully")