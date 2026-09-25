from openpyxl import load_workbook
from openpyxl.styles import Border, Side

wb = load_workbook("company_data.xlsx")
sheet = wb["Employees"]

# Create border style
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

# Apply border to the table
for row in sheet.iter_rows(min_row=1, max_row=6, min_col=1, max_col=5):
    for cell in row:
        cell.border = thin_border

wb.save("company_data.xlsx")

print("Borders applied successfully")