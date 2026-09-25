from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

wb = load_workbook("company_data.xlsx")
sheet = wb["Employees"]

# Merge cells
sheet.merge_cells("A8:E8")

# Add title
sheet["A8"] = "EMPLOYEE DETAILS REPORT"

# Center the title
sheet["A8"].alignment = Alignment(horizontal="center")

# Make the title bold
sheet["A8"].font = Font(bold=True, size=14)

# Save the file
wb.save("company_data.xlsx")

print("Cells merged successfully")