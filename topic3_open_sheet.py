# specific worksheet
from openpyxl import load_workbook
wb=load_workbook("company_data.xlsx")
sheet=wb["Employees"]
print(sheet.title)
print(sheet.max_column)
print(sheet.max_row)