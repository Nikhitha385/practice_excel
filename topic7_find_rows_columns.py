# find the position of data in Excel
from openpyxl import load_workbook
wb = load_workbook("company_data.xlsx")
sheet = wb["Employees"]
# Find the number of rows
print("Rows:", sheet.max_row)
# Find the number of columns
print("Columns:", sheet.max_column)