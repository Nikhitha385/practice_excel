#use a loop to read all employees from the Excel sheet.
from openpyxl import load_workbook
wb = load_workbook("company_data.xlsx")
sheet = wb["Employees"]
for row in sheet.iter_rows(values_only=True):
    print(row)